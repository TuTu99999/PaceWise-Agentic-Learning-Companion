from __future__ import annotations

import re
import shutil
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = ROOT / "outputs" / "course_materials"


COURSES = {
    "database_systems": {
        "code": "DB201",
        "name": "Hệ cơ sở dữ liệu",
        "chapters": [
            (1, "Khái quát về các hệ cơ sở dữ liệu", "dữ liệu, thông tin, hệ quản trị cơ sở dữ liệu", "quản lý dữ liệu cho thư viện trường", "so sánh cách lưu tệp và dùng DBMS; xác định người dùng, giao dịch và rủi ro"),
            (2, "Mô hình cơ sở dữ liệu", "thực thể, thuộc tính, liên kết và lực lượng", "mô hình hóa hệ thống đăng ký câu lạc bộ", "xác định thực thể, khóa, liên kết và chuyển các quy tắc nghiệp vụ thành mô hình ER"),
            (3, "Mô hình cơ sở dữ liệu quan hệ", "quan hệ, bộ, miền, khóa và đại số quan hệ", "thiết kế lược đồ quản lý mượn thiết bị", "chuyển mô hình ER thành lược đồ quan hệ và viết biểu thức đại số quan hệ"),
            (4, "SQL truy vấn dữ liệu", "SELECT, JOIN, GROUP BY, HAVING và truy vấn con", "báo cáo phân tích đơn hàng", "viết truy vấn từ cơ bản đến tổng hợp, xử lý đồng hạng và kiểm tra kết quả"),
            (5, "SQL cập nhật DDL và bảo mật", "DDL, INSERT, UPDATE, DELETE, VIEW và phân quyền", "khởi tạo cơ sở dữ liệu phòng khám", "tạo lược đồ, nạp dữ liệu, viết giao dịch cập nhật và thiết kế quyền truy cập"),
            (6, "Ràng buộc toàn vẹn", "toàn vẹn thực thể, tham chiếu, miền và nghiệp vụ", "kiểm soát dữ liệu đặt phòng", "cài đặt constraint, trigger cần thiết và bộ ca kiểm thử dữ liệu sai"),
            (7, "Phụ thuộc hàm và khóa", "phụ thuộc hàm, bao đóng, phủ tối thiểu và khóa ứng viên", "phân tích lược đồ quản lý đào tạo", "tính bao đóng, tìm khóa, rút gọn tập phụ thuộc và giải thích dư thừa"),
            (8, "Phân tách lược đồ", "phân tách bảo toàn thông tin và bảo toàn phụ thuộc", "tái cấu trúc dữ liệu hóa đơn", "đánh giá lossless join, dependency preservation và đề xuất lược đồ mới"),
            (9, "Chuẩn hóa", "1NF, 2NF, 3NF và BCNF", "chuẩn hóa cơ sở dữ liệu bán hàng", "nhận diện anomaly, chuẩn hóa theo từng bước và kiểm chứng kết quả"),
            (10, "Phụ thuộc đa trị và kết nối", "MVD, 4NF, join dependency và 5NF", "tối ưu lược đồ kỹ năng dự án", "phát hiện phụ thuộc đa trị, phân rã 4NF và kiểm tra phép nối"),
        ],
    },
    "data_structures_and_algorithms": {
        "code": "DSA201",
        "name": "Cấu trúc dữ liệu và giải thuật",
        "chapters": [
            (1, "Tổng quan cấu trúc dữ liệu và giải thuật", "độ phức tạp, Big O và mô hình dữ liệu", "bộ phân tích hiệu năng thuật toán", "phân tích số phép toán và xây bảng so sánh thời gian chạy"),
            (2, "Kỹ thuật tìm kiếm", "tìm kiếm tuần tự, nhị phân và điều kiện áp dụng", "tra cứu hồ sơ sinh viên", "cài đặt hai thuật toán, xử lý dữ liệu trùng và đo số phép so sánh"),
            (3, "Sắp xếp đổi chỗ", "bubble sort, interchange sort và cải tiến dừng sớm", "xếp hạng kết quả thi", "cài đặt, ghi trace từng vòng lặp và so sánh trên ba kiểu dữ liệu"),
            (4, "Sắp xếp chọn và chèn", "selection sort, insertion sort và tính ổn định", "sắp lịch ưu tiên công việc", "cài đặt hai thuật toán, kiểm thử ổn định và chọn thuật toán theo ngữ cảnh"),
            (5, "Sắp xếp trộn và sắp xếp ngoài", "merge sort, chia để trị và external merge", "xử lý tệp log dung lượng lớn", "thiết kế chia run, trộn nhiều lượt và ước lượng I O"),
            (6, "Danh sách đặc", "mảng, danh sách tuần tự và quản lý kích thước", "danh sách sản phẩm trong kho", "cài đặt thêm xóa tìm sửa và xử lý khi mảng đầy"),
            (7, "Danh sách liên kết đơn", "node, con trỏ, chèn xóa và duyệt", "quản lý playlist", "cài đặt danh sách, thao tác tại nhiều vị trí và kiểm tra rò rỉ bộ nhớ"),
            (8, "Danh sách liên kết kép", "liên kết trước sau và duyệt hai chiều", "lịch sử điều hướng trình duyệt", "cài đặt back forward, xóa nhánh lịch sử và kiểm tra biên"),
            (9, "Ngăn xếp và hàng đợi", "stack, queue, deque và ứng dụng", "bộ xử lý biểu thức và hàng chờ dịch vụ", "chuyển biểu thức, tính hậu tố và mô phỏng hàng đợi"),
            (10, "Cây và cây nhị phân", "biểu diễn cây, duyệt preorder inorder postorder", "cây thư mục học liệu", "xây cây, duyệt nhiều thứ tự và tính chiều cao số lá"),
            (11, "Cây nhị phân tìm kiếm", "BST, tìm kiếm, chèn và xóa", "chỉ mục từ điển", "cài đặt ba trường hợp xóa, truy vấn khoảng và đo chiều cao"),
            (12, "Cây AVL và ôn tập", "rotation, hệ số cân bằng và AVL", "chỉ mục sản phẩm cân bằng", "cài đặt bốn phép quay, trace quá trình cân bằng và so sánh với BST"),
        ],
    },
    "object_oriented_programming": {
        "code": "OOPCPP201",
        "name": "Lập trình hướng đối tượng C++",
        "chapters": [
            (1, "Giới thiệu lập trình hướng đối tượng", "đối tượng, lớp, đóng gói và thông điệp", "mô hình hóa hệ thống thuê xe", "nhận diện lớp, trách nhiệm và quan hệ từ yêu cầu nghiệp vụ"),
            (2, "Mở rộng C++ phần 1", "reference, const, namespace và cấp phát động", "bộ quản lý điểm sinh viên", "tổ chức mã nguồn, dùng tham chiếu hợp lý và kiểm soát vùng nhớ"),
            (3, "Mở rộng C++ phần 2", "con trỏ, mảng động, string và xử lý tệp", "trình phân tích nhật ký", "đọc tệp, quản lý bộ nhớ động và tổng hợp thống kê"),
            (4, "Lớp và đối tượng", "thuộc tính, phương thức, phạm vi truy cập và đóng gói", "quản lý tài khoản ngân hàng", "thiết kế lớp, bảo vệ bất biến và viết chương trình minh họa"),
            (5, "Constructor và Destructor", "khởi tạo, hủy, copy constructor và vòng đời", "lớp quản lý bộ đệm", "cài đặt Rule of Three và kiểm thử sao chép sâu"),
            (6, "Thành viên đặc biệt", "static, friend, this và const member", "thống kê đối tượng trong hệ thống", "dùng static để theo dõi, friend có kiểm soát và chuỗi gọi phương thức"),
            (7, "Đa năng hóa toán tử phần 1", "operator số học, so sánh và gán", "thư viện phân số", "nạp chồng toán tử, rút gọn kết quả và xử lý mẫu số không hợp lệ"),
            (8, "Đa năng hóa toán tử phần 2", "operator chỉ mục, nhập xuất và chuyển kiểu", "vector toán học động", "nạp chồng chỉ mục an toàn, stream và phép chuyển đổi"),
            (9, "Tính kế thừa", "lớp cơ sở, lớp dẫn xuất và protected", "hệ thống nhân sự", "thiết kế cây kế thừa, tái sử dụng hành vi và tránh che khuất dữ liệu"),
            (10, "Tính đa hình", "virtual, override, abstract class và dynamic binding", "hệ thống tính phí vận chuyển", "cài đặt interface, đa hình runtime và quản lý đối tượng qua con trỏ cơ sở"),
            (11, "Template và nhập xuất", "function template, class template, stream và serialization", "kho dữ liệu tổng quát", "thiết kế container template, kiểm tra kiểu và lưu tải dữ liệu"),
            (12, "Thiết kế chương trình hướng đối tượng", "SOLID, composition, dependency và kiểm thử", "mini project quản lý thư viện", "phân tích yêu cầu, thiết kế lớp, triển khai theo module và kiểm thử tích hợp"),
        ],
    },
    "web-application-development": {
        "code": "WEB201",
        "name": "Phát triển ứng dụng Web",
        "chapters": [
            (1, "Công nghệ Web và môi trường phát triển", "HTTP, client server, URL và công cụ phát triển", "trang chẩn đoán request response", "quan sát network, giải thích vòng đời request và cấu hình dự án"),
            (2, "HTML và khả năng truy cập", "semantic HTML, form, bảng và accessibility", "website giới thiệu câu lạc bộ", "xây cấu trúc nhiều trang, form hợp lệ và kiểm tra bằng bàn phím"),
            (3, "CSS và thiết kế đáp ứng", "box model, flexbox, grid và media query", "dashboard học tập responsive", "xây layout, hệ màu, trạng thái tương tác và ba breakpoint"),
            (4, "JavaScript trên trình duyệt", "DOM, event, module, fetch và local state", "ứng dụng quản lý công việc", "thêm sửa xóa lọc, kiểm tra dữ liệu và đồng bộ giao diện"),
            (5, "PHP và xử lý phía máy chủ", "request, form, session, validation và template", "cổng đăng ký sự kiện", "xử lý form, session đăng nhập giả lập và phản hồi lỗi an toàn"),
            (6, "PHP kết nối MySQL", "PDO, prepared statement, CRUD và transaction", "mini project quản lý khóa học", "thiết kế database, xây CRUD, tìm kiếm phân trang và giao dịch"),
        ],
    },
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def remove_paragraph_borders(element) -> None:
    p_pr = element.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    p_pr.append(borders)


def add_labelled_paragraph(doc: Document, label: str, content: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(label + " ")
    run.bold = True
    paragraph.add_run(content)


def apply_document_style(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(9)
    remove_paragraph_borders(title._element)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def chapter_questions(topic: str, concepts: str, project: str, work: str, chapter: int) -> dict[str, list[str]]:
    concept_items = [item.strip() for item in concepts.split(",")]
    while len(concept_items) < 4:
        concept_items.append(topic)
    return {
        "foundation": [
            f"Câu 1.1. Giải thích {concept_items[0]} bằng ví dụ tự xây dựng; chỉ ra một trường hợp dùng sai và cách sửa.",
            f"Câu 1.2. So sánh {concept_items[1]} với {concept_items[2]}; nêu tiêu chí lựa chọn trong một tình huống thực tế.",
            f"Câu 1.3. Phân tích một đầu vào hoặc mô hình nhỏ liên quan đến {topic.lower()} và dự đoán kết quả trước khi thực hiện.",
            f"Câu 1.4. Tạo ba ca kiểm thử gồm ca thường, ca biên và ca lỗi cho nội dung {concept_items[3].lower()}.",
        ],
        "practice": [
            f"Yêu cầu 2.1. {work.split(';')[0].strip().capitalize()}; ghi rõ giả định và dữ liệu đầu vào.",
            f"Yêu cầu 2.2. {work.split(';')[-1].strip().capitalize()}; lưu lại bằng chứng kết quả để người khác kiểm tra.",
            "Yêu cầu 2.3. Cài đặt hoặc biểu diễn giải pháp theo quy ước của môn; không bỏ qua xử lý trường hợp biên.",
            "Yêu cầu 2.4. Kiểm thử ít nhất năm trường hợp, giải thích một lỗi tìm được và cách khắc phục.",
        ],
        "project": [
            f"Thành phần 3.1. Phân tích yêu cầu và thiết kế giải pháp cho {project}; lập checklist đầu ra trước khi triển khai.",
            "Thành phần 3.2. Xây dựng phiên bản chạy được gồm dữ liệu mẫu, chức năng chính và xử lý lỗi có thể quan sát.",
            "Thành phần 3.3. Rà soát chất lượng, tối ưu một điểm có căn cứ và viết hướng dẫn chạy hoặc tái hiện kết quả.",
        ],
    }


def build_document(course: dict, chapter: tuple, destination: Path) -> None:
    number, topic, concepts, project, work = chapter
    doc = Document()
    apply_document_style(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run(f"{course['code']} Bài tập Chương {number}")
    remove_paragraph_borders(title._p)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run(topic)
    run.bold = True
    run.font.size = Pt(13)
    intro = doc.add_paragraph(
        "Bộ bài tập gồm câu hỏi nền tảng, bài thực hành nhiều ý và một mini project. "
        "Hãy nộp sản phẩm có thể kiểm tra; không cần ghi ước lượng thời gian vì PaceWise sẽ phân tích từng phần khi lập kế hoạch."
    )
    intro.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=2, cols=3)
    table.autofit = False
    widths = (Cm(3.2), Cm(8.4), Cm(4.8))
    headers = ("Môn học", "Trọng tâm", "Loại nhiệm vụ")
    values = (f"{course['code']} · {course['name']}", concepts, "Cá nhân hoặc cặp đôi")
    for col, width in enumerate(widths):
        for row in table.rows:
            row.cells[col].width = width
    for col, value in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = value
        set_cell_shading(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for col, value in enumerate(values):
        cell = table.cell(1, col)
        cell.text = value
        if col % 2 == 0:
            set_cell_shading(cell, "F3F7FB")
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    heading = doc.add_paragraph(f"Buổi {number} — {topic}", style="Heading 1")
    keep_with_next(heading)
    doc.add_paragraph(
        f"Mục tiêu là vận dụng {concepts} trong một sản phẩm có bằng chứng. Các phần dưới đây được đánh số rõ để Agent có thể tách thành nhiều phiên Do và ước lượng thời lượng theo độ phức tạp."
    )
    questions = chapter_questions(topic, concepts, project, work, number)

    doc.add_paragraph("Câu hỏi nền tảng", style="Heading 2")
    for item in questions["foundation"]:
        add_labelled_paragraph(doc, item.split(". ", 1)[0] + ".", item.split(". ", 1)[1])

    doc.add_paragraph("Thực hành nhiều ý", style="Heading 2")
    for item in questions["practice"]:
        add_labelled_paragraph(doc, item.split(". ", 1)[0] + ".", item.split(". ", 1)[1])

    doc.add_paragraph("Mini project", style="Heading 2")
    project_intro = doc.add_paragraph()
    project_intro.add_run("Bối cảnh. ").bold = True
    project_intro.add_run(f"Xây dựng {project}. Có thể dùng dữ liệu giả lập nhưng phải nêu rõ cấu trúc và giới hạn.")
    for item in questions["project"]:
        add_labelled_paragraph(doc, item.split(". ", 1)[0] + ".", item.split(". ", 1)[1])

    doc.add_paragraph("Sản phẩm nộp", style="Heading 2")
    for text in (
        "Tệp nguồn hoặc mô hình có thể mở và kiểm tra.",
        "Tệp README ngắn mô tả cách chạy, dữ liệu mẫu và quyết định thiết kế chính.",
        "Bằng chứng kiểm thử kèm kết quả mong đợi và kết quả thực tế.",
        "Phần tự đánh giá nêu nội dung đã hoàn thành, nội dung còn thiếu và một cải tiến tiếp theo.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_paragraph("Tiêu chí đánh giá", style="Heading 2")
    doc.add_paragraph(
        "Rubric 10 điểm: Đúng yêu cầu và kiến thức 4; Thiết kế hoặc lập luận 2; Kiểm thử và bằng chứng 2; Chất lượng trình bày và khả năng tái hiện 1; Tự đánh giá trung thực 1."
    )
    doc.add_paragraph(
        "Lưu ý liêm chính học thuật: được dùng AI để giải thích khái niệm, gợi ý bước làm và rà soát; không nộp nguyên lời giải do AI tạo thay cho phần thực hành của người học."
    )

    core = doc.core_properties
    core.title = f"{course['code']} Bài tập Chương {number} {topic}"
    core.subject = "Đề bài dành cho sinh viên và dữ liệu phân rã kế hoạch"
    core.author = "PaceWise Demo Course Team"
    core.keywords = f"{course['code']}, chương {number}, bài tập, mini project"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def safe_name(value: str) -> str:
    value = value.replace("Đ", "D").replace("đ", "d")
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"[^A-Za-z0-9]+", "_", value)
    return value.strip("_") or "chapter"


def main() -> None:
    generated = 0
    for course_key, course in COURSES.items():
        destination_dir = OUTPUT_ROOT / course_key / "exercises" / "by_chapter"
        if destination_dir.exists():
            shutil.rmtree(destination_dir)
        destination_dir.mkdir(parents=True, exist_ok=True)
        for chapter in course["chapters"]:
            number, topic, *_ = chapter
            filename = f"{course['code']}_Bai_tap_Chuong_{number:02d}_{safe_name(topic)}.docx"
            build_document(course, chapter, destination_dir / filename)
            generated += 1
    if generated != 40:
        raise RuntimeError(f"Expected 40 documents, generated {generated}")
    print(f"Generated {generated} chapter assignment documents")


if __name__ == "__main__":
    main()
