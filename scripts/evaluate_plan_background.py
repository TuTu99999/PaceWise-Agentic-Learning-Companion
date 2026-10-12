"""Offline/background evaluation for PaceWise Plan.

This evaluator deliberately bypasses HTTP, authentication, repositories, and
database sessions.  It reads the medium-load demo workbook, builds synthetic
calendar constraints in memory, invokes the Plan decomposition function, and
then runs the canonical deterministic scheduler.

The only side effect is a JSON report under ``outputs/plan-background-eval``.
It never confirms a draft or writes a workspace/calendar record.
"""

# ruff: noqa: E402, I001

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.agents.workflows.plan import plan_node
from src.services.planning import build_plan_draft


DEFAULT_WORKBOOK = (
    ROOT
    / "tests"
    / "fixtures"
    / "planning-personas-medium.xlsx"
)
DEFAULT_REPORT = ROOT / "outputs" / "plan-background-eval" / "report.json"
TIMEZONE_NAME = "Asia/Ho_Chi_Minh"
TIMEZONE = ZoneInfo(TIMEZONE_NAME)
FIXED_NOW = datetime(2026, 10, 11, 13, 30, tzinfo=TIMEZONE)
COURSE_ORDER = ("DB201", "DSA201", "OOPCPP201", "WEB201")
COURSE_IDS = {code: index + 201 for index, code in enumerate(COURSE_ORDER)}
EXPLICIT_ESTIMATES = {
    "DB201": 90,
    "DSA201": 150,
    "OOPCPP201": 120,
    "WEB201": 90,
}


@dataclass(frozen=True)
class WorkbookData:
    lessons: list[dict[str, Any]]
    tasks: list[dict[str, Any]]


def _rows(sheet: Any) -> list[dict[str, Any]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    return [dict(zip(headers, row, strict=True)) for row in values[1:]]


def read_workbook(path: Path) -> WorkbookData:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return WorkbookData(
            lessons=_rows(workbook["Lessons"]),
            tasks=_rows(workbook["Tasks"]),
        )
    finally:
        workbook.close()


def _as_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        parsed = datetime.fromisoformat(str(value))
    return parsed.replace(tzinfo=TIMEZONE) if parsed.tzinfo is None else parsed.astimezone(TIMEZONE)


def _clock(value: Any) -> str:
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    return str(value)[:5]


def _task_id(course_code: str, title: str) -> str:
    folded = re.sub(r"[^a-z0-9]+", "-", title.casefold()).strip("-")
    return f"{course_code}:{folded[:54]}"


def select_assignments(data: WorkbookData) -> list[dict[str, Any]]:
    """Pick the nearest open chapter task for each of the four demo courses."""
    candidates: dict[str, list[dict[str, Any]]] = {code: [] for code in COURSE_ORDER}
    for row in data.tasks:
        course = str(row.get("course_code") or "").strip()
        if course not in candidates or not bool(row.get("published", True)):
            continue
        deadline = _as_datetime(row.get("deadline"))
        if deadline < FIXED_NOW:
            continue
        candidates[course].append({**row, "deadline_value": deadline})

    assignments: list[dict[str, Any]] = []
    for course in COURSE_ORDER:
        if not candidates[course]:
            raise RuntimeError(f"Workbook has no open task for {course} at {FIXED_NOW.isoformat()}")
        row = min(candidates[course], key=lambda item: item["deadline_value"])
        title = str(row.get("title") or "").strip()
        chapter = str(row.get("chapter") or "").strip()
        assignments.append(
            {
                "id": _task_id(course, title),
                "course_id": COURSE_IDS[course],
                "course_code": course,
                "title": title,
                "task_type": str(row.get("kind") or "assignment"),
                "deadline": row["deadline_value"].isoformat(),
                "points_possible": float(row.get("points") or 0),
                "estimated_hours": EXPLICIT_ESTIMATES[course] / 60,
                "knowledge_components": [chapter] if chapter else [],
                "assignment_group": {"name": "Bai tap", "weight": 10},
                "submission": {"workflow_state": "unsubmitted"},
                "priority": "medium",
            }
        )
    return assignments


def lesson_constraints(data: WorkbookData, *, end: date) -> list[dict[str, Any]]:
    slots: list[dict[str, Any]] = []
    for row in data.lessons:
        lesson_day = _as_datetime(row.get("date")).date()
        if FIXED_NOW.date() <= lesson_day <= end:
            slots.append(
                {
                    "date": lesson_day.isoformat(),
                    "weekday": lesson_day.weekday(),
                    "start": _clock(row.get("start_time")),
                    "end": _clock(row.get("end_time")),
                    "category": "class",
                    "source": "workbook",
                    "title": str(row.get("title") or "Buoi hoc"),
                    "courseCode": str(row.get("course_code") or ""),
                }
            )
    return slots


def availability() -> list[dict[str, Any]]:
    return [
        {"weekday": weekday, "enabled": True, "start": "07:00", "end": "23:00"}
        for weekday in range(7)
    ]


def personal_constraints() -> list[dict[str, Any]]:
    return [
        {
            "date": "2026-10-11",
            "weekday": 6,
            "start": "13:30",
            "end": "13:45",
            "category": "personal",
            "title": "An trua va di chuyen (gia lap)",
        },
        {
            "date": "2026-10-11",
            "weekday": 6,
            "start": "18:30",
            "end": "19:00",
            "category": "personal",
            "title": "Viec gia dinh (gia lap)",
        },
        {
            "date": "2026-10-13",
            "weekday": 1,
            "start": "19:00",
            "end": "20:30",
            "category": "personal",
            "title": "Lam them (gia lap)",
        },
        {
            "date": "2026-10-11",
            "weekday": 6,
            "start": "15:00",
            "end": "17:00",
            "category": "self_study",
            "title": "Gio tu hoc uu tien (gia lap)",
        },
    ]


def tasks_input(assignments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "taskId": item["id"],
            "remainingMin": EXPLICIT_ESTIMATES[item["course_code"]],
        }
        for item in assignments
    ]


def base_plan_input(data: WorkbookData, assignments: list[dict[str, Any]]) -> dict[str, Any]:
    dsa = next(item for item in assignments if item["course_code"] == "DSA201")
    return {
        "weekStart": "2026-10-11",
        "horizonEnd": "2026-10-17",
        "tasks": tasks_input(assignments),
        "availability": availability(),
        "unavailable": [
            *lesson_constraints(data, end=date(2026, 10, 17)),
            *personal_constraints(),
        ],
        "preferred": ["morning", "afternoon", "evening"],
        "studyPreferences": {
            "sessionMinutes": 60,
            "breakMinutes": 10,
            "checkpointMinutes": 5,
            "adaptiveSuggestionsEnabled": True,
        },
        "pinnedBlocks": [
            {
                "id": "fake-student-pin-dsa",
                "taskId": dsa["id"],
                "date": "2026-10-11",
                "start": "15:00",
                "end": "16:00",
                "placementSource": "student",
                "checkpointMinutes": 5,
                "activityType": "practice",
            }
        ],
    }


def run_decomposition(
    assignments: list[dict[str, Any]], plan_input: dict[str, Any], *, skip_model: bool
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if skip_model:
        return [], {
            "status": "skipped",
            "elapsedMs": 0,
            "routes": {},
            "qualityPassed": True,
            "checks": {"skipped": True},
        }

    started = time.perf_counter()
    result = plan_node(
        {
            "query": "Chia nho bon bai tap gan nhat, chi dung noi dung chuong cua tung bai.",
            "assignments": assignments,
            "availability": plan_input["availability"],
            "unavailable_slots": plan_input["unavailable"],
            "goals": [
                {
                    "task_id": item["taskId"],
                    "remaining_minutes": item["remainingMin"],
                }
                for item in plan_input["tasks"]
            ],
            "learner_profile": {},
            "mastery_profile": {},
            "messages": [],
        }
    )
    milestones = result.get("proposed_milestones") or []
    response = str(result.get("response") or "")
    model_boundary = (result.get("output") or {}).get("model_output_validation") or {}
    fallback = (
        not milestones
        or bool(model_boundary.get("fallbackApplied"))
        or "fallback" in response.casefold()
        or "chua the" in response.casefold()
    )
    expected_task_ids = {str(item["id"]) for item in assignments}
    milestone_task_ids = {
        str(item.get("task_id") or item.get("taskId") or "")
        for item in milestones
        if str(item.get("task_id") or item.get("taskId") or "")
    }
    unknown_task_ids = sorted(milestone_task_ids - expected_task_ids)
    missing_task_ids = sorted(expected_task_ids - milestone_task_ids)
    milestone_ids = {
        str(item.get("milestone_id") or item.get("id") or "")
        for item in milestones
        if str(item.get("milestone_id") or item.get("id") or "")
    }
    unknown_dependencies = sorted(
        {
            str(dependency)
            for item in milestones
            for dependency in item.get("depends_on") or item.get("dependsOn") or []
            if str(dependency) not in milestone_ids
        }
    )
    allowed_components = {
        str(item["id"]): {str(component).casefold() for component in item.get("knowledge_components") or []}
        for item in assignments
    }
    out_of_scope_components: list[dict[str, str]] = []
    for item in milestones:
        task_id = str(item.get("task_id") or item.get("taskId") or "")
        allowed = allowed_components.get(task_id, set())
        for component in item.get("knowledge_components") or item.get("knowledgeComponents") or []:
            raw_name = (
                component.get("name") or component.get("componentId") or component.get("id")
                if isinstance(component, dict)
                else component
            )
            component_name = str(raw_name or "").strip()
            if component_name and component_name.casefold() not in allowed:
                out_of_scope_components.append({"taskId": task_id, "component": component_name})
    checks = {
        "allSelectedTasksCovered": not missing_task_ids,
        "noUnknownTaskIds": not unknown_task_ids,
        "dependenciesResolve": not unknown_dependencies,
        "knowledgeComponentsStayInChapter": not out_of_scope_components,
        "coarseValidationPassed": bool(((result.get("output") or {}).get("plan_validation") or {}).get("valid")),
    }
    return milestones, {
        "status": "safe_fallback" if fallback else "model_response_parsed",
        "elapsedMs": round((time.perf_counter() - started) * 1000),
        "milestoneCount": len(milestones),
        "coveredTaskIds": sorted(milestone_task_ids),
        "missingTaskIds": missing_task_ids,
        "unknownTaskIds": unknown_task_ids,
        "unknownDependencies": unknown_dependencies,
        "outOfScopeKnowledgeComponents": out_of_scope_components,
        "qualityPassed": all(checks.values()),
        "checks": checks,
        "milestoneSummary": [
            {
                "milestoneId": item.get("milestone_id") or item.get("id"),
                "taskId": item.get("task_id") or item.get("taskId"),
                "title": item.get("title"),
                "estimatedMinutes": item.get("estimated_minutes") or item.get("estimatedMin"),
                "knowledgeComponents": item.get("knowledge_components") or item.get("knowledgeComponents") or [],
                "sourceRefs": item.get("source_refs") or item.get("sourceRefs") or [],
            }
            for item in milestones
        ],
        "routes": (result.get("output") or {}).get("decomposition_routes") or {},
        "modelOutputValidation": model_boundary,
        "validation": (result.get("output") or {}).get("plan_validation") or {},
    }


def _minutes(block: dict[str, Any]) -> int:
    start_hour, start_minute = map(int, str(block["start"]).split(":"))
    end_hour, end_minute = map(int, str(block["end"]).split(":"))
    return end_hour * 60 + end_minute - start_hour * 60 - start_minute


def _overlaps(block: dict[str, Any], slot: dict[str, Any]) -> bool:
    if str(block.get("date"))[:10] != str(slot.get("date"))[:10]:
        return False
    return str(block.get("start")) < str(slot.get("end")) and str(slot.get("start")) < str(block.get("end"))


def verify_draft(
    name: str,
    draft: dict[str, Any],
    assignments: list[dict[str, Any]],
    *,
    expect_feasible: bool,
    require_student_pin: bool,
) -> dict[str, Any]:
    executable = [block for block in draft.get("blocks") or [] if block.get("status") != "history"]
    deadlines = {str(item["id"]): _as_datetime(item["deadline"]) for item in assignments}
    unavailable = [
        slot for slot in (draft.get("input") or {}).get("unavailable") or [] if slot.get("category") != "self_study"
    ]
    assertions: list[dict[str, Any]] = []

    def check(code: str, passed: bool, detail: str) -> None:
        assertions.append({"code": code, "passed": bool(passed), "detail": detail})

    check(
        "EXPECTED_FEASIBILITY",
        bool((draft.get("feasibility") or {}).get("feasible")) is expect_feasible,
        f"expected={expect_feasible}",
    )
    check(
        "QUALITY_HAS_NO_HARD_VIOLATIONS",
        not (draft.get("qualityEvaluation") or {}).get("violations"),
        json.dumps((draft.get("qualityEvaluation") or {}).get("violations") or [], ensure_ascii=False),
    )
    check(
        "NO_BLOCK_AFTER_DEADLINE",
        all(
            datetime.fromisoformat(f"{block['date']}T{block['end']}:00").replace(tzinfo=TIMEZONE)
            <= deadlines[str(block.get("taskId"))]
            for block in executable
        ),
        "All executable blocks end by their task deadline.",
    )
    check(
        "NO_BUSY_OVERLAP",
        not any(_overlaps(block, slot) for block in executable for slot in unavailable),
        "No study block intersects class or synthetic personal busy time.",
    )
    check(
        "SESSION_BOUND_RESPECTED",
        all(int(block.get("workMinutes") or max(0, _minutes(block) - int(block.get("checkpointMinutes") or 0))) <= 60 for block in executable),
        "Work portion of every generated block is at most 60 minutes.",
    )
    check(
        "ONLY_SELECTED_TASKS",
        {str(block.get("taskId")) for block in executable}.issubset(deadlines),
        "No task from another workbook chapter/course leaked into the plan.",
    )
    if require_student_pin:
        check(
            "STUDENT_PIN_PRESERVED",
            any(
                block.get("id") == "fake-student-pin-dsa"
                and block.get("placementSource") == "student"
                and block.get("date") == "2026-10-11"
                and block.get("start") == "15:00"
                and block.get("end") == "16:00"
                for block in executable
            ),
            "The explicit student-positioned study block survives planning/replanning.",
        )
    return {
        "name": name,
        "passed": all(item["passed"] for item in assertions),
        "assertions": assertions,
        "feasibility": draft.get("feasibility") or {},
        "quality": draft.get("qualityEvaluation") or {},
        "priorityOrder": draft.get("priorityOrder") or [],
        "blockCount": len(executable),
        "blocks": executable,
    }


def medium_scenario(
    data: WorkbookData,
    assignments: list[dict[str, Any]],
    milestones: list[dict[str, Any]],
) -> tuple[dict[str, Any], dict[str, Any]]:
    plan_input = base_plan_input(data, assignments)
    draft = build_plan_draft(
        plan_input,
        assignments,
        proposed_milestones=milestones,
        learner_profile={},
        mastery_profile={},
        timezone_name=TIMEZONE_NAME,
        now=FIXED_NOW,
        window_days=7,
    )
    return draft, verify_draft(
        "medium_feasible",
        draft,
        assignments,
        expect_feasible=True,
        require_student_pin=True,
    )


def pressure_scenario(
    data: WorkbookData,
    assignments: list[dict[str, Any]],
    milestones: list[dict[str, Any]],
) -> tuple[dict[str, Any], dict[str, Any]]:
    plan_input = base_plan_input(data, assignments)
    plan_input["pinnedBlocks"] = []
    plan_input["unavailable"] = [
        *[slot for slot in plan_input["unavailable"] if slot.get("category") != "self_study"],
        {
            "date": "2026-10-11",
            "weekday": 6,
            "start": "14:30",
            "end": "22:00",
            "category": "personal",
            "title": "Su co ca nhan keo dai (gia lap)",
        },
    ]
    draft = build_plan_draft(
        plan_input,
        assignments,
        proposed_milestones=milestones,
        learner_profile={},
        mastery_profile={},
        timezone_name=TIMEZONE_NAME,
        now=FIXED_NOW,
        window_days=7,
    )
    result = verify_draft(
        "capacity_pressure",
        draft,
        assignments,
        expect_feasible=False,
        require_student_pin=False,
    )
    result["assertions"].append(
        {
            "code": "SHORTFALL_IS_EXPLICIT",
            "passed": bool((draft.get("feasibility") or {}).get("shortfalls"))
            and int((draft.get("feasibility") or {}).get("deficitMin") or 0) > 0,
            "detail": "Capacity pressure must be reported, not hidden by after-deadline placement.",
        }
    )
    result["passed"] = all(item["passed"] for item in result["assertions"])
    return draft, result


def replan_scenario(
    data: WorkbookData,
    assignments: list[dict[str, Any]],
    milestones: list[dict[str, Any]],
    original_draft: dict[str, Any],
) -> tuple[dict[str, Any], dict[str, Any]]:
    old_agent_blocks = [
        block
        for block in original_draft.get("blocks") or []
        if block.get("placementSource") == "agent" and block.get("status") != "history"
    ]
    if not old_agent_blocks:
        raise RuntimeError("Medium scenario did not produce an agent block for replan testing")
    displaced = old_agent_blocks[0]
    displaced_start_hour, displaced_start_minute = map(int, str(displaced["start"]).split(":"))
    disruption_end = displaced_start_hour * 60 + displaced_start_minute + min(30, _minutes(displaced))
    new_busy = {
        "date": displaced["date"],
        "weekday": date.fromisoformat(displaced["date"]).weekday(),
        "start": displaced["start"],
        "end": f"{disruption_end // 60:02d}:{disruption_end % 60:02d}",
        "category": "personal",
        "title": "Viec dot xuat phat sinh sau khi lap ke hoach (gia lap)",
    }
    plan_input = base_plan_input(data, assignments)
    plan_input["unavailable"] = [*plan_input["unavailable"], new_busy]
    plan_input["pinnedBlocks"] = deepcopy(original_draft.get("blocks") or [])
    draft = build_plan_draft(
        plan_input,
        assignments,
        proposed_milestones=milestones,
        learner_profile={},
        mastery_profile={},
        timezone_name=TIMEZONE_NAME,
        now=FIXED_NOW,
        window_days=7,
    )
    result = verify_draft(
        "replan_after_new_busy_slot",
        draft,
        assignments,
        expect_feasible=True,
        require_student_pin=True,
    )
    new_agent_signatures = {
        (block.get("taskId"), block.get("date"), block.get("start"), block.get("end"))
        for block in draft.get("blocks") or []
        if block.get("placementSource") == "agent" and block.get("status") != "history"
    }
    displaced_signature = (
        displaced.get("taskId"),
        displaced.get("date"),
        displaced.get("start"),
        displaced.get("end"),
    )
    result["assertions"].extend(
        [
            {
                "code": "NEW_BUSY_SLOT_RESPECTED",
                "passed": not any(
                    _overlaps(block, new_busy)
                    for block in draft.get("blocks") or []
                    if block.get("status") != "history"
                ),
                "detail": f"Synthetic disruption: {new_busy}",
            },
            {
                "code": "OLD_AGENT_PLACEMENT_RELEASED",
                "passed": displaced_signature not in new_agent_signatures,
                "detail": "A future agent block is recomputed instead of being treated as a student lock.",
            },
        ]
    )
    result["passed"] = all(item["passed"] for item in result["assertions"])
    result["disruption"] = new_busy
    return draft, result


def compact_scenario(result: dict[str, Any]) -> dict[str, Any]:
    """Keep the JSON useful without duplicating the complete internal draft."""
    return {
        "name": result["name"],
        "passed": result["passed"],
        "assertions": result["assertions"],
        "feasibility": result["feasibility"],
        "quality": {
            key: result["quality"].get(key)
            for key in ("releaseReady", "canConfirm", "score", "checks", "violations", "warnings")
        },
        "priorityOrder": result["priorityOrder"],
        "blockCount": result["blockCount"],
        "blocks": result["blocks"],
        **({"disruption": result["disruption"]} if result.get("disruption") else {}),
    }


def build_report(workbook_path: Path, *, skip_model: bool) -> dict[str, Any]:
    data = read_workbook(workbook_path)
    assignments = select_assignments(data)
    baseline_input = base_plan_input(data, assignments)
    milestones, decomposition = run_decomposition(assignments, baseline_input, skip_model=skip_model)

    medium_draft, medium = medium_scenario(data, assignments, milestones)
    _, pressure = pressure_scenario(data, assignments, milestones)
    _, replan = replan_scenario(data, assignments, milestones, medium_draft)
    scenarios = [medium, pressure, replan]
    return {
        "suite": "pacewise-plan-background-evaluation",
        "passed": all(item["passed"] for item in scenarios) and bool(decomposition.get("qualityPassed")),
        "finalPlanSafetyPassed": all(item["passed"] for item in scenarios),
        "agentDecompositionPassed": bool(decomposition.get("qualityPassed")),
        "executedAt": datetime.now(TIMEZONE).isoformat(),
        "simulationNow": FIXED_NOW.isoformat(),
        "isolation": {
            "httpRequests": 0,
            "databaseWrites": 0,
            "workspaceWrites": 0,
            "calendarWrites": 0,
            "onlyPersistentSideEffect": "This JSON report",
        },
        "workbook": {
            "path": str(workbook_path),
            "lessonRows": len(data.lessons),
            "taskRows": len(data.tasks),
        },
        "selectedAssignments": [
            {
                "id": item["id"],
                "courseCode": item["course_code"],
                "title": item["title"],
                "deadline": item["deadline"],
                "explicitRemainingMin": EXPLICIT_ESTIMATES[item["course_code"]],
                "knowledgeComponents": item["knowledge_components"],
            }
            for item in assignments
        ],
        "decomposition": decomposition,
        "scenarios": [compact_scenario(item) for item in scenarios],
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument(
        "--skip-model",
        action="store_true",
        help="Skip the model decomposition call and evaluate deterministic fallback scheduling only.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    report_path = args.report.resolve()
    report = build_report(workbook_path, skip_model=args.skip_model)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps({
        "passed": report["passed"],
        "report": str(report_path),
        "decomposition": report["decomposition"],
        "scenarios": [
            {
                "name": item["name"],
                "passed": item["passed"],
                "feasible": item["feasibility"].get("feasible"),
                "deficitMin": item["feasibility"].get("deficitMin"),
                "blockCount": item["blockCount"],
            }
            for item in report["scenarios"]
        ],
    }, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
