import base64
import io
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from src.services.teacher_workbook import HEADERS, parse_workbook

COURSES = [{"id": 1, "code": "CS101", "archived": False}, {"id": 2, "code": "OLD", "archived": True}]


def workbook(sheet="Lessons", values=None, headers=None):
    book = Workbook()
    page = book.active
    page.title = sheet
    page.append(headers or HEADERS.get(sheet, ["wrong"]))
    if values is not None:
        page.append(values)
    stream = io.BytesIO()
    book.save(stream)
    return base64.b64encode(stream.getvalue()).decode()


def test_template_round_trip_and_typed_timezone():
    path = Path(__file__).resolve().parents[2] / "frontend/pacewise/public/templates/pacewise-schedule-template.xlsx"
    rows, errors = parse_workbook(base64.b64encode(path.read_bytes()).decode(), COURSES)
    assert not errors and len(rows) == 3
    assert rows[0]["data"]["starts_at"] == "2026-10-12T09:00:00+07:00"
    assert rows[0]["data"]["repeat_count"] == 4
    assert rows[1]["data"]["due_at"] == "2026-10-16T23:59:00+07:00"
    assert rows[1]["data"]["published"] is False
    assert rows[2]["data"]["groups"][0]["weight"] == 100


@pytest.mark.parametrize("code", ["MISSING", "OLD"])
def test_unknown_or_archived_courses_are_not_created_implicitly(code):
    rows, errors = parse_workbook(
        workbook(values=[code, "Class", "2026-10-12", "09:00", "10:00", 1, "offline", "", ""]), COURSES
    )
    assert not rows and "chưa có" in errors[0]["message"]


def test_courses_sheet_defines_new_courses_without_writing_during_parse():
    book = Workbook()
    courses = book.active
    courses.title = "Courses"
    courses.append(HEADERS["Courses"])
    courses.append(["NEW201", "Môn nhập từ Excel", "HK1 2026–2027"])
    lessons = book.create_sheet("Lessons")
    lessons.append(HEADERS["Lessons"])
    lessons.append(["NEW201", "Buổi đầu", "2026-10-12", "09:00", "10:00", 1, "offline", "A101", ""])
    stream = io.BytesIO()
    book.save(stream)

    rows, errors, definitions = parse_workbook(
        base64.b64encode(stream.getvalue()).decode(), COURSES, include_course_definitions=True
    )

    assert not errors
    assert definitions == [
        {"temp_id": -1, "code": "NEW201", "name": "Môn nhập từ Excel", "term": "HK1 2026–2027", "row": 2}
    ]
    assert rows[0]["course_id"] == -1 and rows[0]["course_code"] == "NEW201"


def test_formula_is_not_evaluated_and_error_contains_row():
    _, errors = parse_workbook(
        workbook(
            values=["CS101", '=HYPERLINK("https://example.test")', "2026-10-12", "09:00", "10:00", 1, "offline", "", ""]
        ),
        COURSES,
    )
    assert errors[0]["source"] == "Lessons dòng 2" and "công thức" in errors[0]["message"]


def test_bad_headers_extra_columns_empty_and_unsupported_sheet():
    assert parse_workbook(workbook(headers=["oops"]), COURSES)[1]
    assert parse_workbook(workbook(headers=HEADERS["Lessons"] + ["user_id"]), COURSES)[1]
    assert parse_workbook(workbook(), COURSES)[1]
    assert parse_workbook(workbook(sheet="Macros"), COURSES)[1]


@pytest.mark.parametrize("encoded", ["not base64", base64.b64encode(b"not a zip").decode()])
def test_invalid_files_fail_safely(encoded):
    with pytest.raises(HTTPException) as error:
        parse_workbook(encoded, COURSES)
    assert error.value.status_code == 422


def test_macros_and_external_links_are_rejected_without_loading():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(HTTPException):
        parse_workbook(base64.b64encode(stream.getvalue()).decode(), COURSES)


def test_task_boolean_is_not_guesswork():
    rows, errors = parse_workbook(
        workbook(sheet="Tasks", values=["CS101", "Task", "assignment", datetime(2026, 10, 15, 12), 10, "maybe", "Ch1"]),
        COURSES,
    )
    assert not rows and "TRUE" in errors[0]["message"]


def test_task_group_name_is_imported_and_legacy_task_header_still_works():
    rows, errors = parse_workbook(
        workbook(
            sheet="Tasks",
            values=["CS101", "Task", "assignment", datetime(2026, 10, 15, 12), 10, True, "Ch1", "Bài tập"],
        ),
        COURSES,
    )
    assert not errors and rows[0]["data"]["group_name"] == "Bài tập"

    rows, errors = parse_workbook(
        workbook(
            sheet="Tasks",
            headers=HEADERS["Tasks"][:-1],
            values=["CS101", "Legacy", "assignment", datetime(2026, 10, 15, 12), 10, True, "Ch1"],
        ),
        COURSES,
    )
    assert not errors and rows[0]["data"]["group_name"] == ""


def test_excessive_dimensions_are_rejected():
    book = Workbook()
    page = book.active
    page.title = "Lessons"
    page.append(HEADERS["Lessons"])
    page.cell(10000, 1, "CS101")
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(HTTPException):
        parse_workbook(base64.b64encode(stream.getvalue()).decode(), COURSES)


def test_malformed_streamed_sheet_returns_validation_error():
    source = base64.b64decode(workbook(values=["CS101", "Class", "2026-10-12", "09:00", "10:00", 1, "offline", "", ""]))
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as archive, zipfile.ZipFile(stream, "w") as changed:
        for entry in archive.infolist():
            data = archive.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</sheetData>", b"</broken>")
            changed.writestr(entry.filename, data)
    with pytest.raises(HTTPException) as error:
        parse_workbook(base64.b64encode(stream.getvalue()).decode(), COURSES)
    assert error.value.status_code == 422
