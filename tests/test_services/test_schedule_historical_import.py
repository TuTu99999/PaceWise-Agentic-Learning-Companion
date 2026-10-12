"""Historical timetable import must not manufacture attendance or learning evidence."""

# Imported fixtures are intentionally re-exported for pytest discovery.
# ruff: noqa: F811

import base64
import io
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from src.course_source.db import SessionLocal as CourseSession
from src.course_source.models import (
    Assignment,
    AssignmentGroup,
    CalendarEvent,
    Course,
    CourseLifecycle,
    Enrollment,
)
from src.services.teacher_workbook import HEADERS
from tests.test_services.test_standalone_source import native, sign_in  # noqa: F401

ROOT = "/api/v1/instructor/manage/schedule"


def lesson(start, title="Historical lesson", repeat_count=1):
    return {"course_id": 1, "sheet": "Lessons", "row": 2, "data": {
        "title": title, "starts_at": start.isoformat(),
        "ends_at": (start + timedelta(minutes=90)).isoformat(),
        "repeat_count": repeat_count, "mode": "offline",
    }}


async def preview(client, headers, rows):
    return await client.post(ROOT + "/preview", headers=headers,
                             json={"mode": "append", "course_ids": [1], "rows": rows})


@pytest.mark.asyncio
async def test_historical_and_future_occurrences_import_without_invented_attendance(client, native):
    headers = await sign_in(client)
    start = datetime.now(UTC) - timedelta(days=8)
    response = await preview(client, headers, [lesson(start, repeat_count=3)])
    assert response.status_code == 200, response.text
    draft = response.json()
    assert not draft["errors"] and draft["counts"]["lessons"] == 3
    assert draft["counts"]["historical_lessons"] == 2
    with CourseSession() as db:
        assert not db.scalars(select(CalendarEvent)).all()  # Preview writes no calendar history.
    confirmation = {"ticket": draft["ticket"]}
    first = await client.post(ROOT + "/confirm", headers=headers, json=confirmation)
    assert first.status_code == 200, first.text
    assert (await client.post(ROOT + "/confirm", headers=headers, json=confirmation)).json() == first.json()
    with CourseSession() as db:
        for actor in [1, 2]:
            records = db.scalars(select(CalendarEvent).where(CalendarEvent.user_id == actor)
                                 .order_by(CalendarEvent.starts_at)).all()
            assert len(records) == 3
            assert [event.attendance_status for event in records] == [None, None, "scheduled"]


@pytest.mark.asyncio
async def test_existing_historical_conflict_is_detected(client, native):
    start = datetime.now(UTC) - timedelta(days=30)
    with CourseSession() as db:
        db.add(CalendarEvent(user_id=1, course_id=1, title="Existing history",
                             starts_at=start, ends_at=start + timedelta(minutes=90),
                             event_kind="class", attendance_status="attended"))
        db.commit()
    response = await preview(client, await sign_in(client), [lesson(start)])
    assert response.status_code == 200
    assert any("Trùng" in error["message"] for error in response.json()["errors"])
    assert response.json()["counts"]["historical_lessons"] == 0


@pytest.mark.asyncio
async def test_archived_course_lessons_do_not_block_a_new_timetable(client, native):
    start = datetime.now(UTC) + timedelta(days=30)
    with CourseSession() as db:
        db.add(CourseLifecycle(course_id=2, archived=True))
        db.add(
            CalendarEvent(
                user_id=1,
                course_id=2,
                title="Archived timetable",
                starts_at=start,
                ends_at=start + timedelta(minutes=90),
                event_kind="class_session",
                attendance_status="scheduled",
            )
        )
        db.commit()

    response = await preview(client, await sign_in(client), [lesson(start, title="Replacement timetable")])
    assert response.status_code == 200, response.text
    draft = response.json()
    assert not draft["errors"]
    assert draft["counts"]["lessons"] == 1


@pytest.mark.asyncio
async def test_ongoing_class_and_overdue_task_are_still_rejected(client, native):
    now = datetime.now(UTC)
    rows = [lesson(now - timedelta(minutes=10)),
            {"course_id": 1, "sheet": "Tasks", "row": 2, "data": {
                "title": "Old assignment", "kind": "assignment", "points_possible": 10,
                "due_at": (now - timedelta(days=1)).isoformat(), "published": True,
            }}]
    response = await preview(client, await sign_in(client), rows)
    assert response.status_code == 200, response.text
    errors = response.json()["errors"]
    assert any("đang diễn ra" in error["message"] for error in errors)
    assert any("Deadline phải ở tương lai" in error["message"] for error in errors)


@pytest.mark.asyncio
async def test_excel_preview_creates_declared_course_only_after_confirmation(client, native):
    headers = await sign_in(client)
    future = datetime.now(UTC) + timedelta(days=3)
    book = Workbook()
    courses = book.active
    courses.title = "Courses"
    courses.append(HEADERS["Courses"])
    courses.append(["AUTO401", "Môn tạo từ Excel", "HK1 2026–2027"])
    lessons = book.create_sheet("Lessons")
    lessons.append(HEADERS["Lessons"])
    lessons.append(["AUTO401", "Buổi học thử", future.date(), "09:00", "10:30", 1, "offline", "A101", ""])
    tasks = book.create_sheet("Tasks")
    tasks.append(HEADERS["Tasks"])
    tasks.append(["AUTO401", "Bài tập thử", "assignment", (future + timedelta(days=2)).replace(tzinfo=None), 10, True, "Chương 1", "Bài tập"])
    weights = book.create_sheet("Weights")
    weights.append(HEADERS["Weights"])
    weights.append(["AUTO401", "Bài tập", 100])
    stream = io.BytesIO()
    book.save(stream)

    draft_response = await client.post(
        ROOT + "/import-preview",
        headers=headers,
        json={"mode": "append", "data_base64": base64.b64encode(stream.getvalue()).decode()},
    )
    assert draft_response.status_code == 200, draft_response.text
    draft = draft_response.json()
    assert not draft["errors"]
    assert draft["counts"]["courses"] == 1
    with CourseSession() as db:
        assert not db.scalar(select(Course).where(Course.code == "AUTO401"))

    applied = await client.post(ROOT + "/confirm", headers=headers, json={"ticket": draft["ticket"]})
    assert applied.status_code == 200, applied.text
    assert applied.json()["counts"]["courses"] == 1
    with CourseSession() as db:
        course = db.scalar(select(Course).where(Course.code == "AUTO401"))
        assert course
        assert db.scalar(select(Enrollment).where(Enrollment.course_id == course.id, Enrollment.user_id == 1))
        assert db.scalar(select(CalendarEvent).where(CalendarEvent.course_id == course.id))
        assert db.scalar(select(Assignment).where(Assignment.course_id == course.id))
        assert db.scalar(select(AssignmentGroup).where(AssignmentGroup.course_id == course.id).where(AssignmentGroup.name == "Bài tập"))

    current = (await client.get("/api/v1/auth/me")).json()["user"]
    assert str(applied.json()["created_courses"][0]["id"]) in current["courseIds"]
