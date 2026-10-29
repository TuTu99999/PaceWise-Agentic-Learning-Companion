"""Build a balanced 100-question evaluation set for the four demo courses.

The DB201, DSA201 and WEB201 labels come from the student-facing Quizizz
banks. OOPCPP201 uses a curated set because the workbook currently stored in
that course directory is explicitly an embedded-systems bank, not OOP C++.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

SELECTED_ROWS = (2, 4, 6, 8, 10, 12, 14, 16, 18, 20, 22, 24, 26, 27, 29, 31, 33, 35, 37, 39, 41, 43, 45, 47, 49)

QUIZ_BANKS = {
    "DB201": "information-systems-analysis-design/course/database_systems/quizizz/quizizz_Co_So_Du_Lieu_48_cau.xlsx",
    "DSA201": "data_structures_and_algorithms/quizizz/quizizz_Cau_Truc_Du_Lieu_Va_Giai_Thuat_48_cau.xlsx",
    "WEB201": "web-application-development/quizizz/quizizz_phat_trien_ung_dung_web_48_cau.xlsx",
}

OOP_CASES = (
    ("Nền tảng OOP", "Nhận biết", "Bốn đặc trưng cơ bản của lập trình hướng đối tượng là gì?", "Trừu tượng, đóng gói, kế thừa và đa hình."),
    ("Nền tảng OOP", "Thông hiểu", "Lớp và đối tượng khác nhau như thế nào trong C++?", "Lớp là khuôn mẫu mô tả dữ liệu và hành vi; đối tượng là một thể hiện cụ thể được tạo từ lớp."),
    ("Nền tảng OOP", "Thông hiểu", "Đóng gói dữ liệu trong OOP mang lại lợi ích gì?", "Đóng gói che giấu trạng thái nội bộ và kiểm soát việc truy cập dữ liệu qua giao diện của lớp."),
    ("Mở rộng C++", "Nhận biết", "Namespace trong C++ được dùng để làm gì?", "Namespace gom nhóm định danh và tránh xung đột tên giữa các phần của chương trình."),
    ("Mở rộng C++", "Thông hiểu", "Tham chiếu trong C++ khác con trỏ ở điểm cơ bản nào?", "Tham chiếu là bí danh của một đối tượng và phải được khởi tạo khi khai báo; con trỏ lưu địa chỉ, có thể đổi đối tượng trỏ tới và có thể null."),
    ("Mở rộng C++", "Nhận biết", "Hàm inline trong C++ có ý nghĩa gì?", "Inline đề nghị trình biên dịch thay lời gọi hàm bằng thân hàm tại vị trí gọi để giảm chi phí gọi hàm nhỏ; trình biên dịch có thể từ chối."),
    ("Quản lý bộ nhớ", "Thông hiểu", "Cặp toán tử nào dùng để cấp phát và giải phóng động một đối tượng trong C++?", "Dùng new để cấp phát và delete để giải phóng một đối tượng; với mảng dùng new[] và delete[]."),
    ("Hàm C++", "Nhận biết", "Tham số mặc định của hàm trong C++ được sử dụng khi nào?", "Giá trị mặc định được dùng khi lời gọi hàm bỏ qua đối số tương ứng, và các tham số mặc định phải nằm ở cuối danh sách tham số."),
    ("Lớp và đối tượng", "Nhận biết", "Ba phạm vi truy cập chính của thành viên lớp C++ là gì?", "public, protected và private."),
    ("Lớp và đối tượng", "Thông hiểu", "Con trỏ this trong phương thức thành viên C++ trỏ tới đâu?", "this trỏ tới đối tượng hiện tại đang gọi phương thức thành viên không tĩnh."),
    ("Constructor và Destructor", "Nhận biết", "Constructor có nhiệm vụ chính gì?", "Constructor khởi tạo trạng thái của đối tượng khi đối tượng được tạo và có cùng tên với lớp, không có kiểu trả về."),
    ("Constructor và Destructor", "Nhận biết", "Destructor trong C++ được gọi khi nào và dùng để làm gì?", "Destructor được gọi khi đối tượng bị hủy để giải phóng tài nguyên; tên của nó là tên lớp có dấu ngã ở trước."),
    ("Constructor và Destructor", "Thông hiểu", "Copy constructor được dùng trong trường hợp nào?", "Copy constructor tạo một đối tượng mới từ một đối tượng cùng lớp, chẳng hạn khi khởi tạo bằng đối tượng khác hoặc truyền/trả về theo giá trị."),
    ("Thành viên đặc biệt", "Thông hiểu", "Thành viên static của lớp khác thành viên thông thường như thế nào?", "Thành viên static thuộc về lớp và được dùng chung cho mọi đối tượng, thay vì mỗi đối tượng có một bản riêng."),
    ("Thành viên đặc biệt", "Nhận biết", "Hàm friend có quyền truy cập gì đối với một lớp C++?", "Hàm hoặc lớp friend có thể truy cập các thành viên private và protected của lớp đã khai báo friendship."),
    ("Đa năng hóa toán tử", "Thông hiểu", "Mục đích của operator overloading trong C++ là gì?", "Nạp chồng toán tử định nghĩa cách một toán tử sẵn có hoạt động với đối tượng của lớp do người dùng tạo."),
    ("Đa năng hóa toán tử", "Thông hiểu", "Nạp chồng toán tử có thay đổi độ ưu tiên và số toán hạng của toán tử không?", "Không. Nạp chồng không thay đổi độ ưu tiên, tính kết hợp hoặc số lượng toán hạng vốn có của toán tử."),
    ("Kế thừa", "Nhận biết", "Quan hệ IS-A trong OOP thường được biểu diễn bằng cơ chế nào?", "Quan hệ IS-A được biểu diễn bằng kế thừa, trong đó lớp dẫn xuất là một dạng chuyên biệt của lớp cơ sở."),
    ("Kế thừa", "Thông hiểu", "Thứ tự gọi constructor và destructor trong quan hệ kế thừa là gì?", "Constructor lớp cơ sở chạy trước constructor lớp dẫn xuất; khi hủy, destructor lớp dẫn xuất chạy trước destructor lớp cơ sở."),
    ("Kế thừa", "Thông hiểu", "Kế thừa public ảnh hưởng thế nào đến thành viên public và protected của lớp cơ sở?", "Với kế thừa public, thành viên public của lớp cơ sở vẫn public và thành viên protected vẫn protected trong lớp dẫn xuất."),
    ("Đa hình", "Nhận biết", "Hàm virtual trong C++ hỗ trợ cơ chế nào?", "Hàm virtual hỗ trợ liên kết động và đa hình thời gian chạy khi gọi qua con trỏ hoặc tham chiếu lớp cơ sở."),
    ("Đa hình", "Nhận biết", "Hàm thuần ảo và lớp trừu tượng là gì?", "Hàm thuần ảo được khai báo với = 0; lớp có ít nhất một hàm thuần ảo là lớp trừu tượng và không thể tạo đối tượng trực tiếp."),
    ("Đa hình", "Thông hiểu", "Vì sao destructor của lớp cơ sở đa hình thường phải là virtual?", "Destructor virtual bảo đảm xóa đối tượng dẫn xuất qua con trỏ lớp cơ sở sẽ gọi đầy đủ destructor dẫn xuất rồi destructor cơ sở."),
    ("Template", "Nhận biết", "Function template trong C++ dùng để làm gì?", "Function template mô tả một thuật toán tổng quát có thể được biên dịch cho nhiều kiểu dữ liệu."),
    ("Template và nhập xuất", "Thông hiểu", "Class template và luồng tệp trong C++ lần lượt giải quyết vấn đề gì?", "Class template tạo kiểu lớp tổng quát theo tham số kiểu; ifstream và ofstream hỗ trợ đọc và ghi dữ liệu từ tệp."),
)


def _quiz_cases(root: Path, course_code: str, relative_path: str) -> list[dict]:
    workbook = load_workbook(root / relative_path, read_only=True, data_only=True)
    sheet = workbook["Ngan hang chi tiet"]
    cases = []
    for ordinal, row_number in enumerate(SELECTED_ROWS, 1):
        values = [sheet.cell(row_number, column).value for column in range(1, 11)]
        answer_letter = str(values[9]).strip().upper()
        answer_index = ord(answer_letter) - ord("A")
        options = values[5:9]
        if answer_index not in range(4) or not options[answer_index]:
            raise ValueError(f"Invalid answer at {course_code} row {row_number}")
        cases.append(
            {
                "case_id": f"{course_code}-{ordinal:02d}",
                "course_code": course_code,
                "topic": str(values[2]),
                "level": str(values[3]),
                "user_input": str(values[4]).strip(),
                "reference": str(options[answer_index]).strip(),
                "label_source": f"quizizz:{Path(relative_path).name}:row:{row_number}",
            }
        )
    return cases


def build(material_root: Path) -> list[dict]:
    cases = []
    for course_code in ("DB201", "DSA201"):
        cases.extend(_quiz_cases(material_root, course_code, QUIZ_BANKS[course_code]))
    cases.extend(
        {
            "case_id": f"OOPCPP201-{index:02d}",
            "course_code": "OOPCPP201",
            "topic": topic,
            "level": level,
            "user_input": question,
            "reference": reference,
            "label_source": "curated_from_student_slides",
        }
        for index, (topic, level, question, reference) in enumerate(OOP_CASES, 1)
    )
    cases.extend(_quiz_cases(material_root, "WEB201", QUIZ_BANKS["WEB201"]))
    counts = Counter(case["course_code"] for case in cases)
    if counts != {"DB201": 25, "DSA201": 25, "OOPCPP201": 25, "WEB201": 25}:
        raise ValueError(f"Unbalanced evaluation set: {counts}")
    normalized = [" ".join(case["user_input"].casefold().split()) for case in cases]
    if len(set(normalized)) != 100:
        raise ValueError("Evaluation questions must be unique")
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--material-root", type=Path, default=Path("outputs/course_materials"))
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    cases = build(args.material_root)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        "\n".join(json.dumps(case, ensure_ascii=False) for case in cases) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"sample_count": len(cases), "courses": Counter(c["course_code"] for c in cases)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
