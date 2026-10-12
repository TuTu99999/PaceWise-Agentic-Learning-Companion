"""Bounded, deterministic XLSX parsing; no macros, links, formulas or model calls."""

import base64
import io
import re
import zipfile
from datetime import date, datetime, time
from xml.etree.ElementTree import ParseError
from zoneinfo import ZoneInfo

from fastapi import HTTPException

HEADERS = {
    "Courses": ["course_code", "course_name", "term"],
    "Lessons": [
        "course_code",
        "title",
        "date",
        "start_time",
        "end_time",
        "repeat_count",
        "mode",
        "location",
        "meeting_url",
    ],
    "Tasks": ["course_code", "title", "kind", "deadline", "points", "published", "chapter", "group_name"],
    "Weights": ["course_code", "group_name", "weight"],
}
LOCAL = ZoneInfo("Asia/Ho_Chi_Minh")


def local_instant(value):
    if isinstance(value, datetime):
        result = value
    else:
        result = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return (result if result.tzinfo else result.replace(tzinfo=LOCAL)).isoformat()


def boolean(value):
    if value in (True, 1, "true", "TRUE"):
        return True
    if value in (None, False, 0, "false", "FALSE", ""):
        return False
    raise ValueError("published dùng TRUE hoặc FALSE")


def parse_workbook(encoded, courses, *, include_course_definitions=False):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    errors, rows = [], []
    try:
        content = base64.b64decode(encoded, validate=True)
        if len(content) > 5 * 1024 * 1024:
            raise ValueError("File vượt 5 MB")
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 1000 or sum(item.file_size for item in infos) > 30 * 1024 * 1024:
                raise ValueError("File giải nén quá lớn")
            if any("vbaProject" in item.filename or "externalLinks/" in item.filename for item in infos):
                raise ValueError("Không nhận macro hoặc liên kết workbook ngoài")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (ValueError, zipfile.BadZipFile, KeyError, OSError, ParseError, InvalidFileException, RuntimeError) as exc:
        raise HTTPException(422, "Chỉ nhận file .xlsx hợp lệ, tối đa 5 MB, không macro/liên kết ngoài.") from exc
    codes = {item["code"].casefold(): int(item["id"]) for item in courses if not item["archived"]}
    archived_codes = {item["code"].casefold() for item in courses if item["archived"]}
    new_courses = []
    weight_groups = {}
    total = 0
    try:
        course_sheet = workbook["Courses"] if "Courses" in workbook.sheetnames else None
        if course_sheet is not None:
            if course_sheet.max_row is None or course_sheet.max_column is None:
                course_sheet.calculate_dimension(force=True)
            if course_sheet.max_row > 101 or course_sheet.max_column > 20:
                raise HTTPException(422, "Sheet Courses tối đa 100 dòng và 20 cột.")
            expected = HEADERS["Courses"]
            entries = course_sheet.iter_rows(max_row=101, max_col=max(len(expected), course_sheet.max_column))
            header = next(entries, ())
            header_values = [cell.value for cell in header]
            if header_values[: len(expected)] != expected or any(
                value is not None for value in header_values[len(expected) :]
            ):
                errors.append({"source": "Courses", "message": "Cột không khớp file mẫu: " + ", ".join(expected)})
            else:
                declared = set()
                for index, cells in enumerate(entries, 2):
                    values = [cell.value for cell in cells]
                    if not any(value is not None for value in values):
                        continue
                    total += 1
                    if total > 500:
                        raise HTTPException(422, "Tổng cộng tối đa 500 dòng mỗi lần nhập.")
                    try:
                        if any(cell.data_type in {"f", "e"} for cell in cells):
                            raise ValueError("Không nhận công thức hoặc ô lỗi; nhập giá trị trực tiếp")
                        if any(cell.value is not None for cell in cells[len(expected) :]):
                            raise ValueError("Có dữ liệu ngoài các cột file mẫu")
                        value = dict(zip(expected, values))
                        code = str(value["course_code"] or "").strip()
                        name = str(value["course_name"] or "").strip()
                        term = str(value["term"] or "Demo").strip()
                        key = code.casefold()
                        if not code or len(code) > 64 or not re.fullmatch(r"[\w.-]+", code):
                            raise ValueError("Mã môn chỉ gồm chữ, số, dấu chấm, gạch ngang hoặc gạch dưới; tối đa 64 ký tự")
                        if not name or len(name) > 255:
                            raise ValueError("Tên môn bắt buộc và tối đa 255 ký tự")
                        if not term or len(term) > 128:
                            raise ValueError("Học kỳ bắt buộc và tối đa 128 ký tự")
                        if key in declared:
                            raise ValueError("Mã môn bị lặp trong sheet Courses")
                        declared.add(key)
                        if key in archived_codes:
                            raise ValueError("Môn đã bị xóa/lưu trữ. Khôi phục môn trước khi nhập")
                        if key not in codes:
                            temporary_id = -(len(new_courses) + 1)
                            codes[key] = temporary_id
                            new_courses.append(
                                {
                                    "temp_id": temporary_id,
                                    "code": code,
                                    "name": name,
                                    "term": term,
                                    "row": index,
                                }
                            )
                    except (TypeError, ValueError) as exc:
                        errors.append({"source": f"Courses dòng {index}", "message": str(exc)[:500]})

        for sheet in workbook:
            if sheet.title == "Courses":
                continue
            if sheet.title not in HEADERS:
                errors.append(
                    {"source": sheet.title, "message": "Sheet không được hỗ trợ; dùng Courses, Lessons, Tasks, Weights."}
                )
                continue
            if sheet.max_row is None or sheet.max_column is None:
                # Some standards-compliant exporters omit the dimension element.
                sheet.calculate_dimension(force=True)
            if sheet.max_row > 501 or sheet.max_column > 20:
                raise HTTPException(422, "Mỗi sheet tối đa 500 dòng và 20 cột.")
            expected = HEADERS[sheet.title]
            entries = sheet.iter_rows(max_row=501, max_col=max(len(expected), sheet.max_column))
            header = next(entries, ())
            header_values = [cell.value for cell in header]
            legacy_task_headers = expected[:-1] if sheet.title == "Tasks" else []
            current_headers = header_values[: len(expected)]
            has_expected = current_headers == expected and not any(
                value is not None for value in header_values[len(expected) :]
            )
            has_legacy_task_headers = legacy_task_headers and (
                header_values[: len(legacy_task_headers)] == legacy_task_headers
                and not any(value is not None for value in header_values[len(legacy_task_headers) :])
            )
            if not has_expected and not has_legacy_task_headers:
                errors.append({"source": sheet.title, "message": "Cột không khớp file mẫu: " + ", ".join(expected)})
                continue
            active_headers = expected if has_expected else legacy_task_headers
            for index, cells in enumerate(entries, 2):
                values = [cell.value for cell in cells]
                if not any(value is not None for value in values):
                    continue
                total += 1
                if total > 500:
                    raise HTTPException(422, "Tổng cộng tối đa 500 dòng mỗi lần nhập.")
                try:
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValueError("Không nhận công thức hoặc ô lỗi; nhập giá trị trực tiếp")
                    if any(cell.value is not None for cell in cells[len(active_headers) :]):
                        raise ValueError("Có dữ liệu ngoài các cột file mẫu")
                    value = dict(zip(active_headers, values))
                    course_id = codes.get(str(value["course_code"]).strip().casefold())
                    if course_id is None:
                        raise ValueError(
                            "Mã môn chưa có hoặc bạn không quản lý môn này. Khai báo môn trong sheet Courses hoặc tạo môn trước"
                        )
                    if sheet.title == "Lessons":
                        day = value["date"]
                        if isinstance(day, datetime):
                            day = day.date()
                        if not isinstance(day, date):
                            day = date.fromisoformat(str(day))
                        start = value["start_time"]
                        end = value["end_time"]
                        start = start if isinstance(start, time) else time.fromisoformat(str(start))
                        end = end if isinstance(end, time) else time.fromisoformat(str(end))
                        data = {
                            "title": value["title"],
                            "starts_at": datetime.combine(day, start, LOCAL).isoformat(),
                            "ends_at": datetime.combine(day, end, LOCAL).isoformat(),
                            "repeat_count": value["repeat_count"] or 1,
                            "mode": value["mode"] or "offline",
                            "location": value["location"] or "",
                            "meeting_url": value["meeting_url"] or "",
                        }
                    elif sheet.title == "Tasks":
                        data = {
                            "title": value["title"],
                            "kind": value["kind"] or "assignment",
                            "due_at": local_instant(value["deadline"]),
                            "points_possible": value["points"],
                            "published": boolean(value["published"]),
                            "chapter": value["chapter"] or "",
                            "group_name": str(value.get("group_name") or "").strip(),
                        }
                    else:
                        weight_groups.setdefault(course_id, []).append(
                            {"name": value["group_name"], "weight": value["weight"]}
                        )
                        continue
                    rows.append(
                        {
                            "course_id": course_id,
                            "course_code": str(value["course_code"]).strip(),
                            "sheet": sheet.title,
                            "row": index,
                            "data": data,
                        }
                    )
                except (TypeError, ValueError) as exc:
                    errors.append({"source": f"{sheet.title} dòng {index}", "message": str(exc)[:500]})
        for course_id, groups in weight_groups.items():
            course_code = next((code for code, identifier in codes.items() if identifier == course_id), "")
            rows.append(
                {
                    "course_id": course_id,
                    "course_code": course_code.upper(),
                    "sheet": "Weights",
                    "row": 2,
                    "data": {"groups": groups},
                }
            )
    except (zipfile.BadZipFile, KeyError, OSError, ParseError, InvalidFileException, RuntimeError, ValueError) as exc:
        # Read-only worksheets stream their XML lazily; malformed rows can fail after load_workbook.
        raise HTTPException(422, "Nội dung workbook bị hỏng; xuất lại file .xlsx theo mẫu rồi thử lại.") from exc
    finally:
        workbook.close()
    if not rows and not errors:
        errors.append({"source": "File", "message": "File chưa có dữ liệu."})
    result = (rows, errors, new_courses)
    return result if include_course_definitions else result[:2]
